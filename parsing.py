"""Text extraction for the supported upload formats (PDF, DOCX, TXT, Markdown)."""
from __future__ import annotations

import io
import os

from app.core.errors import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md"})

Page = tuple[int | None, str]


def _read_pdf(data: bytes) -> list[Page]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentParsingError("Encrypted PDFs are not supported.")
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    except DocumentParsingError:
        raise
    except Exception as exc:  # pypdf raises many different error types
        raise DocumentParsingError(f"Could not read PDF: {exc}") from exc


def _read_docx(data: bytes) -> list[Page]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return [(None, "\n\n".join(blocks))]
    except Exception as exc:
        raise DocumentParsingError(f"Could not read DOCX: {exc}") from exc


def extract_pages(filename: str, data: bytes) -> list[Page]:
    """Return ``[(page_number | None, text), ...]`` with empty pages removed."""
    extension = os.path.splitext(filename.lower())[1]
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or 'unknown'}'. Supported types: {supported}."
        )

    if extension == ".pdf":
        pages = _read_pdf(data)
    elif extension == ".docx":
        pages = _read_docx(data)
    else:
        pages = [(None, data.decode("utf-8", errors="replace"))]

    # PostgreSQL text columns cannot hold NUL bytes, which some PDFs produce.
    cleaned = [(n, t.replace("\x00", "").strip()) for n, t in pages]
    cleaned = [(n, t) for n, t in cleaned if t]
    if not cleaned:
        raise EmptyDocumentError(
            "No extractable text found (scanned or image-only documents need OCR)."
        )
    return cleaned
