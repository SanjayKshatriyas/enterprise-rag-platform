import io

from app.core.errors import (
    DocumentParsingError,
    EmptyDocumentError,
    UnsupportedFileTypeError,
)
from app.core.parsing import extract_pages


def expect(exc_type, fn):
    try:
        fn()
    except exc_type:
        return
    raise AssertionError(f"expected {exc_type.__name__}")


def test_text_and_markdown():
    assert extract_pages("n.txt", b"hello  world") == [(None, "hello  world")]
    assert extract_pages("README.MD", "# Title\ntext".encode()) == [(None, "# Title\ntext")]


def test_unsupported_extension():
    expect(UnsupportedFileTypeError, lambda: extract_pages("virus.exe", b"x"))
    expect(UnsupportedFileTypeError, lambda: extract_pages("noextension", b"x"))


def test_empty_and_nul_only_documents():
    expect(EmptyDocumentError, lambda: extract_pages("a.txt", b"   \n  "))
    expect(EmptyDocumentError, lambda: extract_pages("a.txt", b"\x00\x00"))


def test_nul_bytes_are_stripped():
    assert extract_pages("a.txt", b"ab\x00cd") == [(None, "abcd")]


def test_corrupt_files_raise_parsing_error():
    expect(DocumentParsingError, lambda: extract_pages("bad.pdf", b"not a pdf"))
    expect(DocumentParsingError, lambda: extract_pages("bad.docx", b"not a docx"))


def test_docx_paragraphs_and_tables():
    import docx

    document = docx.Document()
    document.add_paragraph("Quarterly security review.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Owner"
    table.rows[0].cells[1].text = "Security team"
    buffer = io.BytesIO()
    document.save(buffer)

    ((page, text),) = extract_pages("report.docx", buffer.getvalue())
    assert page is None
    assert "Quarterly security review." in text
    assert "Owner | Security team" in text


def test_pdf_pages_are_numbered():
    try:
        from reportlab.pdfgen import canvas
    except ImportError:  # reportlab is a dev-only dependency
        return
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 750, "First page about vacation policy.")
    pdf.showPage()
    pdf.drawString(72, 750, "Second page about security policy.")
    pdf.showPage()
    pdf.save()

    pages = extract_pages("policy.pdf", buffer.getvalue())
    assert [n for n, _ in pages] == [1, 2]
    assert "vacation" in pages[0][1] and "security" in pages[1][1]
